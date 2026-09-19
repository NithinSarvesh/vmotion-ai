import os
from pptx import Presentation
from docx import Document

pptx_path = 'VMotion_AI_DA1_Presentation_FINAL.pptx'
docx_path = 'VMotion_AI_DA1_Project_Report_FINAL.docx'

print(f'PPTX size: {os.path.getsize(pptx_path):,} bytes')
print(f'DOCX size: {os.path.getsize(docx_path):,} bytes')

# Audit PPTX
prs = Presentation(pptx_path)
pptx_text = []
for slide_idx, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame:
            pptx_text.append(shape.text_frame.text)
full_pptx = '\n'.join(pptx_text)

# Audit DOCX
doc = Document(docx_path)
docx_text = []
for p in doc.paragraphs:
    docx_text.append(p.text)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            docx_text.append(cell.text)
full_docx = '\n'.join(docx_text)

print('\n--- INSTITUTION AUDIT ---')
bad_terms = ['VIT Vellore', 'VIT, Vellore', 'Vellore campus', 'Vellore Campus', 'campus of Vellore']
for term in bad_terms:
    count_pptx = full_pptx.count(term)
    count_docx = full_docx.count(term)
    print(f'Term "{term}": PPTX={count_pptx}, DOCX={count_docx}')

# Check for proper campus naming
print('"VIT Chennai" occurrences: PPTX=', full_pptx.count('VIT Chennai'), 'DOCX=', full_docx.count('VIT Chennai'))
print('"Chennai" occurrences: PPTX=', full_pptx.count('Chennai'), 'DOCX=', full_docx.count('Chennai'))

# Check for any rogue "Vellore" mentions
# Note that "Vellore Institute of Technology" is fine as long as campus is Chennai
import re
vellore_instances_pptx = [line for line in full_pptx.split('\n') if 'vellore' in line.lower()]
vellore_instances_docx = [line for line in full_docx.split('\n') if 'vellore' in line.lower()]
print(f'\nTotal lines containing "Vellore": PPTX={len(vellore_instances_pptx)}, DOCX={len(vellore_instances_docx)}')
for line in vellore_instances_pptx:
    print(f'  PPTX: {line.strip()[:80]}')
for line in vellore_instances_docx:
    print(f'  DOCX: {line.strip()[:80]}')

# Check for banned AI phrases
banned_phrases = [
    'academic ground rules',
    'catastrophic RL hallucinations',
    'statistical black boxes',
    'formal invariants',
    'do not invent',
    'as an AI',
    'TODO',
    'FIXME',
    'Note:',
    'generator'
]
print('\n--- LANGUAGE & METADATA AUDIT ---')
for phrase in banned_phrases:
    c_p = full_pptx.lower().count(phrase.lower())
    c_d = full_docx.lower().count(phrase.lower())
    if c_p > 0 or c_d > 0:
        print(f'WARNING: "{phrase}" found! PPTX={c_p}, DOCX={c_d}')
    else:
        print(f'Clean: "{phrase}"')

print('\nAudit complete.')
